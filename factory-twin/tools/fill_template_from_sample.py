"""통합 테스트용: 샘플 공장 JSON으로 수집 템플릿을 '현장에서 채운 것처럼' 채운다 → templates_to_json → factory_sim → validate 체인 검증.
실제 방문 데이터가 아니라 변환기 코드 경로를 검증하기 위한 것."""
import json, sys, pathlib, random, datetime
from openpyxl import load_workbook
ROOT = pathlib.Path(__file__).resolve().parent.parent
src = ROOT / "capture/templates/현장수집_템플릿.xlsx"; dst = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/sample_filled_template.xlsx"
f = json.loads((ROOT / "data/sample_factory.json").read_text(encoding="utf-8"))
wb = load_workbook(src); rng = random.Random(3)
ws = wb["0_공장기본"]
for row in ws.iter_rows(min_row=3):
    k = row[0].value
    vals = {"공장명": "샘플 공장(템플릿 테스트)", "건물 가로 x (m)": f["site"]["width"], "건물 세로 y (m)": f["site"]["depth"], "계약전력 (kW)": 300, "평균 단가 (원/kWh)": 165, "기본부하 (kW)": 12, "주 근무일수": 5}
    if k in vals: row[1].value = vals[k]
for z in f["site"]["zones"]:
    xs = [p[0] for p in z["polygon"]]; ys = [p[1] for p in z["polygon"]]
    wb["1_구역"].append([z["id"], z["name"], z["type"], min(xs), min(ys), max(xs), max(ys), ""])
hdr = [c.value for c in wb["2_기계대장"][2]]
for s in f["stations"]:
    r = {h: "" for h in hdr}
    r.update({"기계ID": s["id"], "기계명": s["name"], hdr[2]: s["process"], "구역ID": s.get("zone", ""), "중심 x": s["pos"][0], "중심 y": s["pos"][1], "길이(x,m)": s["size"][0], "깊이(y,m)": s["size"][1], "높이(m)": s["size"][2],
              "대수(capacity)": s.get("capacity", 1), "작업자 수": s.get("workers_required", 1), "정격 kW(명판)": s.get("rated_kw", 0), "집진 연결(ID)": ",".join(s.get("ancillary", [])),
              "입력버퍼 x": s["in_pos"][0], "입력버퍼 y": s["in_pos"][1], "출력버퍼 x": s["out_pos"][0], "출력버퍼 y": s["out_pos"][1], "버퍼 용량(배치/대차 수)": s.get("buffer_batches", 4),
              "셋업 시간(분)": s.get("setup_min", 0), "MTBF(시간)": (s.get("mtbf_min", 0) or 0) / 60, "MTTR(분)": s.get("mttr_min", 0), "불량률(%)": (s.get("scrap_rate", 0) or 0) * 100,
              "소정지 빈도(회/시간)": 1 if s["process"] == "edge" else 0, "이동 등급(A자유/B리깅/C고정)": "B",
              "이송속도 v_f(m/min, 엣지·샌더·런닝쏘)": (12 if s["id"] == "EB1" else 8) if s["process"] == "edge" else "", "부품 간격 G(m, 엣지)": 0.45 if s["process"] == "edge" else ""})
    wb["2_기계대장"].append([r[h] for h in hdr])
for a in f.get("ancillary", []):
    wb["3_부대설비"].append([a["id"], a["name"], "", a["pos"][0], a["pos"][1], a["size"][0], a["size"][1], a["size"][2], a["rated_kw"], a["mode"], a.get("load_factor", 0.7), "", "N", "", "", "", ""])
# 사이클타임: 스테이션당 12회, 가공+핸들링 2행/측정, 엣지는 변 수 드라이버
hdr4 = [c.value for c in wb["4_공정_사이클타임"][2]]
for s in f["stations"]:
    if s["process"] in ("raw", "fg", "dry"): continue
    for i in range(12):
        edges = rng.choice([1, 2, 3, 4]); holes = rng.choice([4, 6, 8, 10, 14])
        if s["process"] == "edge": edge_len = round(edges * rng.uniform(0.4, 1.2), 2); vf = 12 if s["id"] == "EB1" else 8; base = (edge_len + edges * 0.45) / vf + 0.08 + 0.15 * edges  # 이송 + γ0 + γ1·N_pass
        elif s["process"] == "bore": base = 0.5 + 0.035 * holes
        elif s["process"] in ("assemble", "pack"): base = rng.uniform(5, 20)
        else: base = rng.uniform(0.6, 1.2)
        mach = base * 0.7 * rng.uniform(0.9, 1.1); hand = base * 0.3 * rng.uniform(0.8, 1.2)
        for kind, sec in (("가공", mach * 60), ("핸들링", hand * 60)):
            r = {h: "" for h in hdr4}; r.update({"기계ID": s["id"], "제품ID": "P_DESK", "측정#": i + 1, "엣지 변 수": edges, "엣지 길이 합(m)": (edge_len if s["process"] == "edge" else ""), "홀 수": holes, "소요(초)": round(sec, 1), hdr4[10]: kind, "평정(100=정상 속도)": 100, "첫 사이클(버림 Y)": "Y" if i == 0 else ""})
            wb["4_공정_사이클타임"].append([r[h] for h in hdr4])
for sid, m in [("CUT1", 3), ("EB1", 8), ("EB2", 5), ("BOR2", 9)]:
    wb["5_셋업_체인지오버"].append([sid, "화이트/18", "오크/18", "엣지테이프", "", "", m, "", "", "", 4, 1, "", ""])
for w in f["workers"]:
    wb["6_작업자"].append([w["id"], w.get("name", ""), w["role"], w.get("station") or "", ",".join(w.get("skills", [])), "정규", w.get("monthly_wage_krw", 0), "", "", ""])
hdr7 = [c.value for c in wb["7_제품_BOM"][2]]
for p in f["products"]:
    for part in p["parts"]:
        r = {h: "" for h in hdr7}
        r.update({"제품ID": p["id"], "제품명": p["name"], "원장 색상": p.get("color", ""), "원장 두께(mm)": p.get("thickness", 18), "월 수량": 100, hdr7[5]: f["orders"]["batch_size"][p["id"]], hdr7[6]: ">".join(p["routing"]),
                  "부품명": part["id"], "수량/제품": part["qty"], "가로 w(m)": part["w"], "세로 d(m)": part["d"], "두께 t(m)": part.get("t", 0.018), "엣지 변 수": part.get("n_edges", 0), "홀 수": part.get("holes", 0), "도장(Y/N)": "Y" if part.get("finish") else "N",
                  "조립 시간(분/제품)": 15, "포장 시간(분/제품)": 5, "판매가(원)": p.get("price_krw", 0)})
        wb["7_제품_BOM"].append([r[h] for h in hdr7])
d0 = datetime.date(2026, 9, 1)
for i in range(60):
    od = d0 + datetime.timedelta(days=i // 3)
    if od.weekday() >= 5: continue
    pid = rng.choice(list(f["orders"]["mix"])); lead = rng.randint(5, 12)
    wb["8_주문이력"].append([od, od + datetime.timedelta(days=lead), pid, f["orders"]["batch_size"][pid], od + datetime.timedelta(days=lead - 1), 0, ""])
wb["16_의사결정규칙"].append(["투입", "시공일이 정해지면", "시공일 근무일 기준 D-3 아침에 재단 투입", "급한 건은 당일", "반장", ""])
wb["16_의사결정규칙"].append(["배치구성", "같은 날 같은 색상·두께 주문이 여럿이면", "색상·두께로 묶어 재단 후 현장별 분류", "", "반장", ""])
wb["18_검증지표"]["C3"] = 1500; wb["18_검증지표"]["C4"] = 850; wb["18_검증지표"]["C5"] = 0.9; wb["18_검증지표"]["C6"] = 30; wb["18_검증지표"]["C7"] = "CUT2"; wb["18_검증지표"]["C8"] = 0.6
wb["20_외주경로"].append(["도어", "P_CAB,P_WARD", "○○도어", "주문 접수 시", 3, 5, 9, "조립", "조립대 옆 적치", 40, 6, ""])
wb["23_운반_출하대장"].append(["대차", 6, 30, 1.2, "운반", "차면", 1.5, 0.4, "", ""])
wb["23_운반_출하대장"].append(["지게차", 1, 1, 2.0, "지게차 기사", "호출", 2, "", "", ""])
wb["23_운반_출하대장"].append(["탑차", 2, 3, 40, "시공", "", "", "", 4, ""])
wb.save(dst); print("wrote", dst)
