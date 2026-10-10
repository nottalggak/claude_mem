"""시트 13(배치추적, 현장 2주 양식) → trace_validate / rule_fit 입력(worklog.csv, releases.csv).
시트 13 열: 배치번호(스티커), 주문/제품ID, 수량, 공정/기계ID, 도착 시각, 착수 시각, 완료 시각, 출발 시각, 셋업 발생(분), 작업자, 대기 사유, 운반 수단, 부품 수(입/출), 함께 처리한 배치, 외주 반출/반입 시각, 사진, 비고
시각은 'YYYY-MM-DD HH:MM' 또는 'MM/DD HH:MM' 또는 'HH:MM'(앞 행의 날짜 이어받음). Day0 = 기록 첫 날(월요일 정렬: 시뮬 달력과 요일을 맞추기 위해 첫 날의 요일만큼 오프셋).
수량×BOM 부품 수 = n_parts(factory.json 필요), 주문일·시공일은 시트 8(주문이력)에서 배치의 주문/제품ID 로 조인(없으면 투입=첫 도착, 납기 없음).
사용: python3 tools/sheet13_to_worklog.py 채움본.xlsx factory.json worklog.csv releases.csv"""
import sys, csv, json, pathlib, datetime as dt, re
import openpyxl


def parse_t(v, last_date):
    if v in (None, ""): return None, last_date
    if isinstance(v, dt.datetime): return v, v.date()
    if isinstance(v, dt.time): return (dt.datetime.combine(last_date, v) if last_date else None), last_date
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M"):
        try: t = dt.datetime.strptime(s, fmt); return t, t.date()
        except ValueError: pass
    m = re.match(r"^(\d{1,2})/(\d{1,2}) (\d{1,2}):(\d{2})$", s)  # MM/DD HH:MM — 연도는 앞 행에서
    if m:
        y = (last_date or dt.date.today()).year; t = dt.datetime(y, int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4))); return t, t.date()
    m = re.match(r"^(\d{1,2}):(\d{2})$", s)
    if m and last_date: return dt.datetime.combine(last_date, dt.time(int(m.group(1)), int(m.group(2)))), last_date
    return None, last_date


def main():
    xlsx, fjson, out_wl, out_rel = sys.argv[1:5]
    f = json.loads(pathlib.Path(fjson).read_text(encoding="utf-8")); ppu = {p["id"]: sum(x.get("qty", 1) for x in p["parts"]) for p in f["products"]}
    wb = openpyxl.load_workbook(xlsx, data_only=True); ws = wb["13_배치추적"]
    hdr = [c.value for c in ws[2]]; rows = [dict(zip(hdr, [c.value for c in r])) for r in ws.iter_rows(min_row=3) if any(c.value not in (None, "") for c in r)]
    last = None; recs = []
    for r in rows:
        arr, last = parse_t(r.get("도착 시각"), last); st, last = parse_t(r.get("착수 시각"), last); en, last = parse_t(r.get("완료 시각"), last)
        if not (r.get("배치번호(스티커)") and r.get("공정/기계ID") and arr and st and en): continue
        recs.append({"batch": str(r["배치번호(스티커)"]), "product": str(r.get("주문/제품ID") or ""), "qty": int(float(r.get("수량") or 1)), "station": str(r["공정/기계ID"]), "arrive": arr, "start": st, "end": en, "setup": float(r.get("셋업 발생(분, 없으면 0)") or 0)})
    if not recs: print("유효 행 없음(배치번호·기계ID·도착·착수·완료 필요)"); return
    d0 = min(x["arrive"] for x in recs).date(); off = d0.weekday()  # 시뮬 Day0 = 월요일
    tomin = lambda t: ((t.date() - d0).days + off) * 1440 + t.hour * 60 + t.minute + t.second / 60
    # 제품ID 정규화: 주문ID 가 적혀 있으면 시트 8 로 제품·주문일·납기 조인
    orders = {}
    if "8_주문이력" in wb.sheetnames:
        h8 = [c.value for c in wb["8_주문이력"][2]]
        for r in wb["8_주문이력"].iter_rows(min_row=3):
            d = dict(zip(h8, [c.value for c in r]))
            if d.get("제품ID"): orders.setdefault(str(d.get("비고") or ""), d)  # 비고 열에 주문ID 를 적는 관행 — 없으면 제품ID 로만
    with open(out_wl, "w", newline="", encoding="utf-8") as fp:
        w = csv.DictWriter(fp, fieldnames=["batch", "product", "n_parts", "station", "arrive", "start", "end", "setup"]); w.writeheader()
        for x in recs:
            pid = x["product"] if x["product"] in ppu else (orders.get(x["product"], {}).get("제품ID") or x["product"])
            w.writerow({"batch": x["batch"], "product": pid, "n_parts": x["qty"] * ppu.get(pid, 1), "station": x["station"], "arrive": round(tomin(x["arrive"]), 2), "start": round(tomin(x["start"]), 2), "end": round(tomin(x["end"]), 2), "setup": x["setup"]})
    by = {}
    for x in recs: by.setdefault(x["batch"], []).append(x)
    with open(out_rel, "w", newline="", encoding="utf-8") as fp:
        w = csv.DictWriter(fp, fieldnames=["batch", "product", "qty", "release", "due", "order"]); w.writeheader()
        for b, xs in by.items():
            first = min(xs, key=lambda x: x["arrive"]); pid = first["product"] if first["product"] in ppu else (orders.get(first["product"], {}).get("제품ID") or first["product"])
            od = orders.get(first["product"], {}); due = od.get("납기일"); order = od.get("주문일")
            dmin = round(tomin(due if isinstance(due, dt.datetime) else dt.datetime.combine(due, dt.time(17, 0))), 2) if isinstance(due, (dt.datetime, dt.date)) else ""
            omin = round(tomin(order if isinstance(order, dt.datetime) else dt.datetime.combine(order, dt.time(9, 0))), 2) if isinstance(order, (dt.datetime, dt.date)) else ""
            w.writerow({"batch": b, "product": pid, "qty": first["qty"], "release": round(tomin(first["arrive"]), 2), "due": dmin, "order": omin})
    print(f"wrote {out_wl} ({len(recs)} 단계), {out_rel} ({len(by)} 배치); Day0={d0}(요일 오프셋 {off}). 배치 ≥20·공정당 ≥15 이어야 구조 검증이 의미 있다(연구 18 §6).")


if __name__ == "__main__":
    main()
